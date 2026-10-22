"""Shared XLSX export helper used by all four form-type apps.

One row per filing, columns standardized across forms so users can compare:
issuer, filing date, amount raised, amount target/max, industry, contacts."""

from io import BytesIO

from django.http import HttpResponse


def csv_safe_cell(value):
    """Neutralize spreadsheet expressions in text without changing numbers."""
    if isinstance(value, str) and (
        value.lstrip().startswith(("=", "+", "-", "@"))
        or value.startswith(("\t", "\r", "\n"))
    ):
        return "'" + value
    return value


def finish_export(request, response, token):
    """Claim after generation; never return a file if another request won."""
    if token is not None and not token.consume():
        return HttpResponse(
            "This export credit was already used by another request. Please retry with another credit.",
            status=409, content_type="text/plain",
        )
    response["Cache-Control"] = "private, no-store"
    return response


def xlsx_response(filename: str, rows: list[list], headers: list[str]) -> HttpResponse:
    """Build a single-sheet xlsx and wrap it in an HttpResponse."""
    # Import inside the function so the module imports cheaply elsewhere.
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Filings"

    bold = Font(bold=True)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = bold
        cell.alignment = Alignment(horizontal="left")

    for row_index, row in enumerate(rows, start=2):
        ws.append(row)
        # Explicit string cells prevent formula/error interpretation by Excel.
        for column_index, value in enumerate(row, start=1):
            if isinstance(value, str):
                ws.cell(row=row_index, column=column_index).data_type = "s"

    # Reasonable column widths; xlsx auto-fit is expensive, set generous defaults.
    widths = {1: 42, 2: 36, 3: 13, 4: 18, 5: 18, 6: 24, 7: 60}
    for col_idx, width in widths.items():
        ws.column_dimensions[chr(64 + col_idx)].width = width

    ws.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    resp = HttpResponse(
        buf.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
    )
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    return resp


def _related_person_names(filing) -> str:
    """Comma-join the names from a Filing's related_persons (Form D)."""
    try:
        return ", ".join(rp.name for rp in filing.related_persons.all() if rp.name)[:1000]
    except Exception:
        return ""
